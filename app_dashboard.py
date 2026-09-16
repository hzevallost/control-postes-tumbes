# -*- coding: utf-8 -*-
"""
Dashboard de Control de Postes Observados - v3.5
Mejoras principales:
- Google Sheets como fuente de datos.
- Google Drive automático para fotos por código/nombre, enlace o ID.
- Carga/reemplazo de fotos ANTES y DESPUÉS desde la propia app.
- Compatible con carpetas privadas compartidas con una Service Account.
- Filtros, búsqueda, KPIs, gráficos, mapa opcional, control de calidad.
- Exportación a Excel y PDF.
- Preparado para Streamlit Community Cloud.
"""

import io
import os
import re
import hmac
import mimetypes
import unicodedata
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pandas as pd
import streamlit as st
import plotly.express as px
import plotly.graph_objects as go

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Image as RLImage,
    PageBreak,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.utils import ImageReader

# Dependencias Google Drive (el dashboard sigue funcionando sin ellas,
# salvo la lectura automática de carpetas privadas).
try:
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload
    GOOGLE_DRIVE_LIBS_OK = True
except Exception:
    GOOGLE_DRIVE_LIBS_OK = False


# -----------------------------------------------------------------------------
# CONFIGURACIÓN GENERAL
# -----------------------------------------------------------------------------
st.set_page_config(
    page_title="Control de Izaje de Postes",
    page_icon="🏗️",
    layout="wide",
)

st.markdown(
    """
    <style>
        .metric-card {
            background-color: #f8f9fa;
            border: 1px solid #e9ecef;
            padding: 15px;
            border-radius: 10px;
            text-align: center;
            box-shadow: 0 2px 4px rgba(0,0,0,0.05);
            min-height: 105px;
        }
        .metric-title {
            font-size: 12px;
            color: #6c757d;
            font-weight: 700;
            text-transform: uppercase;
        }
        .metric-value {
            font-size: 25px;
            color: #212529;
            font-weight: bold;
        }
        .metric-sub {
            font-size: 11px;
            color: #6c757d;
            margin-top: 3px;
        }
        .block-container {
            padding-top: 1.6rem;
            padding-bottom: 2rem;
        }
        [data-testid="stImage"] {
            overflow: visible !important;
        }
        [data-testid="stImage"] img {
            padding: 4px 0px;
            border-radius: 8px;
        }
        .centered-subheader {
            text-align: center;
            font-weight: bold;
            font-size: 20px;
            color: #212529;
            margin-bottom: 5px;
            display: flex;
            align-items: center;
            justify-content: center;
            gap: 8px;
        }
        .photo-box {
            border: 1px solid #dee2e6;
            border-radius: 10px;
            padding: 10px;
            background: #ffffff;
        }
        .small-muted {
            color: #6c757d;
            font-size: 12px;
        }
    </style>
    """,
    unsafe_allow_html=True,
)

DEFAULT_SHEET_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1HTEq01G5xgyMCNrocYvOeKIXTV0xhsKg/export?format=csv"
)

ESTADOS_BASE = ["PENDIENTE", "ATENDIDO", "CONFORME"]
COLOR_ESTADOS = {
    "PENDIENTE": "#d9534f",
    "ATENDIDO": "#ffc107",
    "CONFORME": "#28a745",
}

IMAGE_MIME_PREFIX = "image/"
FOLDER_MIME = "application/vnd.google-apps.folder"
DRIVE_SCOPE = "https://www.googleapis.com/auth/drive"


# -----------------------------------------------------------------------------
# HELPERS DE CONFIGURACIÓN / SECRETOS
# -----------------------------------------------------------------------------
def _secrets_section(name: str) -> dict:
    try:
        if name in st.secrets:
            return dict(st.secrets[name])
    except Exception:
        pass
    return {}


def _to_list(value) -> List[str]:
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        return [str(x).strip() for x in value if str(x).strip()]
    value = str(value).strip()
    if not value:
        return []
    # Permite coma o punto y coma en configuración local.
    return [x.strip() for x in re.split(r"[;,]", value) if x.strip()]


def get_sheet_url() -> str:
    cfg = _secrets_section("app")
    return str(cfg.get("sheet_url", DEFAULT_SHEET_URL)).strip()


def get_drive_folder_config() -> Dict[str, List[str]]:
    cfg = _secrets_section("drive")
    root = _to_list(cfg.get("root_folder_ids"))
    antes = _to_list(cfg.get("antes_folder_ids")) or root
    despues = _to_list(cfg.get("despues_folder_ids")) or root
    return {
        "root": root,
        "antes": antes,
        "despues": despues,
    }


def has_service_account_secrets() -> bool:
    try:
        return "google_service_account" in st.secrets
    except Exception:
        return False


def get_upload_pin() -> str:
    """PIN opcional para proteger las cargas desde una app compartida."""
    cfg = _secrets_section("app")
    return str(cfg.get("upload_pin", "")).strip()


# -----------------------------------------------------------------------------
# GOOGLE SHEETS
# -----------------------------------------------------------------------------
@st.cache_data(ttl=60, show_spinner=False)
def cargar_datos_gsheets(url: str) -> Tuple[Optional[pd.DataFrame], Optional[str]]:
    try:
        df = pd.read_csv(url, header=None)
        df = df.dropna(how="all")
        if len(df) > 0:
            df.columns = [f"COL_{i}" for i in range(df.shape[1])]
        return df, None
    except Exception as e:
        return None, str(e)


def preparar_dataframe(df_raw: pd.DataFrame) -> pd.DataFrame:
    """Mantiene la lógica del proyecto original y detecta columnas nuevas por encabezado."""
    df = df_raw.copy()

    # Detecta encabezados LATITUD / LONGITUD antes de filtrar las filas de datos.
    # Esto hace que funcione aunque las coordenadas se agreguen al final de Google Sheets.
    encabezados_detectados = {}
    for idx_col in range(df.shape[1]):
        muestra = (
            df.iloc[:10, idx_col]
            .astype(str)
            .str.upper()
            .str.strip()
        )
        if muestra.isin(["LATITUD", "LATITUDE", "LAT"]).any():
            encabezados_detectados[idx_col] = "LATITUD"
        elif muestra.isin(["LONGITUD", "LONGITUDE", "LON", "LONG"]).any():
            encabezados_detectados[idx_col] = "LONGITUD"

    # Busca la columna de ESTADO por contenido.
    col_estado_idx = None
    for col in df.columns:
        valores_str = df[col].astype(str).str.upper().str.strip()
        if valores_str.str.contains(r"PENDIENTE|ATENDIDO|CONFORME|\bOK\b", regex=True).any():
            col_estado_idx = col
            break

    if col_estado_idx is not None:
        df[col_estado_idx] = df[col_estado_idx].astype(str).str.upper().str.strip()
        df[col_estado_idx] = df[col_estado_idx].replace({"OK": "ATENDIDO"})
        # Conserva los tres estados originales. Si luego agregas otros estados,
        # simplemente añádelos en esta lista o elimina este filtro.
        df = df[df[col_estado_idx].isin(ESTADOS_BASE)]

    df = df.reset_index(drop=True)
    df.index = df.index + 1

    # Nombres por posición, compatibles con tu archivo actual.
    if df.shape[1] >= 7:
        base_names = [
            "ZONA",
            "N° POSTE",
            "TIPO / ALTURA",
            "TERRENO",
            "JUSTIFICACIÓN",
            "OBSERVACIÓN / ACCIÓN",
            "ESTADO",
        ]

        extras = []
        if df.shape[1] >= 9:
            extras.extend(["FOTO ANTES", "FOTO DESPUES"])
            extras.extend([f"EXTRA_{i}" for i in range(9, df.shape[1])])
        elif df.shape[1] == 8:
            extras.append("EXTRA_7")

        nombres = (base_names + extras)[: df.shape[1]]

        # Si Google Sheets contiene encabezados LATITUD/LONGITUD, se respetan
        # independientemente de la posición exacta de las columnas.
        for idx_col, nombre_detectado in encabezados_detectados.items():
            if idx_col < len(nombres):
                nombres[idx_col] = nombre_detectado

        # Compatibilidad con la estructura actual: H/I fotos y J/K coordenadas.
        # Solo se aplica si no se detectó explícitamente el encabezado.
        if df.shape[1] >= 10 and "LATITUD" not in nombres:
            nombres[9] = "LATITUD"
        if df.shape[1] >= 11 and "LONGITUD" not in nombres:
            nombres[10] = "LONGITUD"

        df.columns = nombres

    return df


# -----------------------------------------------------------------------------
# GOOGLE DRIVE AUTOMÁTICO
# -----------------------------------------------------------------------------
@st.cache_resource(show_spinner=False)
def get_drive_service():
    if not GOOGLE_DRIVE_LIBS_OK:
        raise RuntimeError(
            "Faltan librerías de Google Drive. Instala google-api-python-client y google-auth."
        )
    if not has_service_account_secrets():
        raise RuntimeError(
            "No se encontró [google_service_account] en los secretos de Streamlit."
        )

    info = dict(st.secrets["google_service_account"])
    creds = service_account.Credentials.from_service_account_info(
        info,
        scopes=[DRIVE_SCOPE],
    )
    return build("drive", "v3", credentials=creds, cache_discovery=False)


def extract_drive_file_id(value: str) -> Optional[str]:
    """Extrae un fileId desde enlaces típicos de Google Drive o acepta un ID puro."""
    if value is None:
        return None
    s = str(value).strip()
    if not s or s.lower() == "nan":
        return None

    patterns = [
        r"/file/d/([a-zA-Z0-9_-]+)",
        r"/d/([a-zA-Z0-9_-]+)",
        r"[?&]id=([a-zA-Z0-9_-]+)",
    ]
    for p in patterns:
        m = re.search(p, s)
        if m:
            return m.group(1)

    # Los IDs reales de Drive suelen ser bastante largos. Evita confundir 116_A con un ID.
    if re.fullmatch(r"[a-zA-Z0-9_-]{20,}", s):
        return s
    return None


def normalize_text_token(value: str) -> str:
    """Normaliza textos para comparar zonas/carpetas sin depender de tildes o signos."""
    if value is None:
        return ""
    s = str(value).strip()
    if not s or s.lower() == "nan":
        return ""
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"[^A-Z0-9]+", "", s.upper())


def normalize_photo_key(value: str) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    if not s or s.lower() == "nan":
        return ""
    stem = Path(s).stem if "/" not in s and "\\" not in s else s
    return normalize_text_token(stem)


def normalize_poste_code(value: str) -> str:
    """Evita que un número leído como 116.0 termine buscando 116.0_A."""
    if value is None:
        return ""
    s = str(value).strip()
    if not s or s.lower() == "nan":
        return ""
    if re.fullmatch(r"[-+]?\d+\.0+", s):
        s = s.split(".", 1)[0]
    return s


def _drive_list_children(service, folder_id: str) -> List[dict]:
    children = []
    page_token = None
    while True:
        response = (
            service.files()
            .list(
                q=f"'{folder_id}' in parents and trashed = false",
                spaces="drive",
                fields=(
                    "nextPageToken, files("
                    "id,name,mimeType,parents,webViewLink,thumbnailLink,modifiedTime,size)"
                ),
                pageToken=page_token,
                pageSize=1000,
                supportsAllDrives=True,
                includeItemsFromAllDrives=True,
            )
            .execute()
        )
        children.extend(response.get("files", []))
        page_token = response.get("nextPageToken")
        if not page_token:
            break
    return children


@st.cache_data(ttl=300, show_spinner=False)
def indexar_drive(folder_ids_tuple: Tuple[str, ...]) -> Tuple[List[dict], Optional[str]]:
    """Indexa recursivamente Drive y conserva la ruta relativa de cada archivo."""
    folder_ids = [x for x in folder_ids_tuple if x]
    if not folder_ids:
        return [], "No hay carpetas de Google Drive configuradas."

    try:
        service = get_drive_service()
        queue = [(folder_id, "") for folder_id in dict.fromkeys(folder_ids)]
        visited = set()
        files = []

        while queue:
            folder_id, relative_folder = queue.pop(0)
            if folder_id in visited:
                continue
            visited.add(folder_id)

            for item in _drive_list_children(service, folder_id):
                if item.get("mimeType") == FOLDER_MIME:
                    child_path = f"{relative_folder}/{item.get('name', '')}".strip("/")
                    queue.append((item["id"], child_path))
                else:
                    item = dict(item)
                    item["root_scanned"] = folder_id
                    item["relative_folder"] = relative_folder
                    item["folder_keys"] = [
                        normalize_text_token(part)
                        for part in relative_folder.split("/")
                        if normalize_text_token(part)
                    ]
                    item["key_name"] = normalize_photo_key(item.get("name", ""))
                    item["key_stem"] = normalize_photo_key(Path(item.get("name", "")).stem)
                    files.append(item)

        return files, None
    except Exception as e:
        return [], str(e)


@st.cache_data(ttl=300, show_spinner=False)
def get_drive_metadata(file_id: str) -> Tuple[Optional[dict], Optional[str]]:
    try:
        service = get_drive_service()
        item = (
            service.files()
            .get(
                fileId=file_id,
                fields="id,name,mimeType,webViewLink,thumbnailLink,modifiedTime,size",
                supportsAllDrives=True,
            )
            .execute()
        )
        return item, None
    except Exception as e:
        return None, str(e)


@st.cache_data(ttl=300, show_spinner=False)
def descargar_archivo_drive(file_id: str) -> Tuple[Optional[bytes], Optional[str], Optional[str]]:
    """Descarga bytes de un archivo blob de Drive mediante la Service Account."""
    try:
        service = get_drive_service()
        meta = (
            service.files()
            .get(
                fileId=file_id,
                fields="id,name,mimeType",
                supportsAllDrives=True,
            )
            .execute()
        )
        request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
        fh = io.BytesIO()
        downloader = MediaIoBaseDownload(fh, request, chunksize=1024 * 1024)
        done = False
        while not done:
            _, done = downloader.next_chunk()
        fh.seek(0)
        return fh.getvalue(), meta.get("mimeType"), meta.get("name")
    except Exception as e:
        return None, None, str(e)


def buscar_carpeta_zona(zona: str, root_folder_ids: List[str]) -> Tuple[Optional[dict], Optional[str]]:
    """Busca recursivamente la carpeta de una zona dentro de las raíces configuradas."""
    zona_key = normalize_text_token(zona)
    if not zona_key:
        return None, "Zona vacía o inválida."
    if not root_folder_ids:
        return None, "No hay una carpeta raíz de Google Drive configurada."

    try:
        service = get_drive_service()
        queue = [(fid, "") for fid in dict.fromkeys(root_folder_ids) if fid]
        visited = set()

        while queue:
            folder_id, rel = queue.pop(0)
            if folder_id in visited:
                continue
            visited.add(folder_id)

            for item in _drive_list_children(service, folder_id):
                if item.get("mimeType") != FOLDER_MIME:
                    continue
                nombre = str(item.get("name", ""))
                ruta = f"{rel}/{nombre}".strip("/")
                if normalize_text_token(nombre) == zona_key:
                    return {"id": item["id"], "name": nombre, "relative_folder": ruta}, None
                queue.append((item["id"], ruta))

        return None, f"No se encontró la carpeta de la zona '{zona}'."
    except Exception as e:
        return None, str(e)


def _extension_imagen(uploaded_file) -> str:
    """Devuelve una extensión segura y compatible para la evidencia cargada."""
    suffix = Path(getattr(uploaded_file, "name", "")).suffix.lower()
    permitidas = {".jpg", ".jpeg", ".png", ".webp"}
    if suffix in permitidas:
        return suffix

    mime = str(getattr(uploaded_file, "type", "") or "").lower()
    por_mime = {
        "image/jpeg": ".jpg",
        "image/png": ".png",
        "image/webp": ".webp",
    }
    return por_mime.get(mime, ".jpg")


def subir_o_reemplazar_foto_poste(
    zona: str,
    poste: str,
    tipo: str,
    uploaded_file,
    root_folder_ids: List[str],
) -> Tuple[Optional[dict], Optional[str], Optional[str]]:
    """
    Sube una foto a la carpeta de la zona. Si ya existe POSTE_A/POSTE_D,
    actualiza el archivo más reciente para evitar duplicados nuevos.

    Retorna: (metadata, acción, error).
    """
    if uploaded_file is None:
        return None, None, "No se seleccionó ningún archivo."

    zona_folder, err = buscar_carpeta_zona(zona, root_folder_ids)
    if err or not zona_folder:
        return None, None, err or "No se pudo resolver la carpeta de zona."

    poste_code = normalize_poste_code(poste)
    if not poste_code:
        return None, None, "Número de poste inválido."

    sufijo = "A" if normalize_text_token(tipo).startswith("A") else "D"
    ext = _extension_imagen(uploaded_file)
    target_name = f"{poste_code}_{sufijo}{ext}"
    target_key = normalize_photo_key(f"{poste_code}_{sufijo}")
    mime = str(getattr(uploaded_file, "type", "") or "").strip()
    if not mime.startswith("image/"):
        mime = mimetypes.guess_type(target_name)[0] or "image/jpeg"

    try:
        data = uploaded_file.getvalue()
        if not data:
            return None, None, "El archivo seleccionado está vacío."

        service = get_drive_service()
        children = _drive_list_children(service, zona_folder["id"])
        existentes = [
            x for x in children
            if str(x.get("mimeType", "")).startswith(IMAGE_MIME_PREFIX)
            and normalize_photo_key(Path(str(x.get("name", ""))).stem) == target_key
        ]
        existentes.sort(key=lambda x: str(x.get("modifiedTime", "")), reverse=True)

        media = MediaIoBaseUpload(io.BytesIO(data), mimetype=mime, resumable=False)

        if existentes:
            elegido = existentes[0]
            item = (
                service.files()
                .update(
                    fileId=elegido["id"],
                    body={"name": target_name},
                    media_body=media,
                    fields="id,name,mimeType,parents,webViewLink,modifiedTime,size",
                    supportsAllDrives=True,
                )
                .execute()
            )
            accion = "reemplazada"
            if len(existentes) > 1:
                accion += f" (se detectaron {len(existentes)} duplicados previos; se actualizó el más reciente)"
        else:
            item = (
                service.files()
                .create(
                    body={"name": target_name, "parents": [zona_folder["id"]]},
                    media_body=media,
                    fields="id,name,mimeType,parents,webViewLink,modifiedTime,size",
                    supportsAllDrives=True,
                )
                .execute()
            )
            accion = "cargada"

        item = dict(item)
        item["relative_folder"] = zona_folder.get("relative_folder", zona_folder.get("name", ""))
        return item, accion, None
    except Exception as e:
        msg = str(e)
        if "403" in msg or "insufficient" in msg.lower() or "permission" in msg.lower():
            msg = (
                "Google Drive rechazó la escritura. Verifica que la cuenta de servicio "
                "tenga permiso de EDITOR sobre la carpeta principal control-postes-tumbes. "
                f"Detalle: {e}"
            )
        return None, None, msg


def limpiar_cache_drive():
    """Fuerza que las fotos nuevas aparezcan inmediatamente tras una carga."""
    try:
        indexar_drive.clear()
    except Exception:
        pass
    try:
        get_drive_metadata.clear()
    except Exception:
        pass
    try:
        descargar_archivo_drive.clear()
    except Exception:
        pass


def resolver_foto(value: str, indexed_files: List[dict]) -> Tuple[Optional[dict], str]:
    """
    Prioridad:
    1) URL o ID directo de Drive.
    2) Coincidencia exacta por nombre/código.
    3) Coincidencia parcial única.
    """
    if value is None or str(value).strip() == "" or str(value).lower() == "nan":
        return None, "vacío"

    s = str(value).strip()
    direct_id = extract_drive_file_id(s)
    if direct_id:
        meta, err = get_drive_metadata(direct_id)
        if meta:
            return meta, "id/enlace"
        return None, f"ID no accesible: {err}"

    key = normalize_photo_key(s)
    if not key:
        return None, "código inválido"

    # Exacta por stem/nombre normalizado.
    exact = [
        f for f in indexed_files
        if f.get("key_stem") == key or f.get("key_name") == key
    ]
    if exact:
        # Prioriza imágenes.
        exact.sort(key=lambda x: 0 if str(x.get("mimeType", "")).startswith(IMAGE_MIME_PREFIX) else 1)
        return exact[0], "nombre exacto"

    # Parcial si solo existe una coincidencia clara.
    partial = [
        f for f in indexed_files
        if key in f.get("key_stem", "") or f.get("key_stem", "") in key
    ]
    partial = [f for f in partial if f.get("key_stem")]
    if len(partial) == 1:
        return partial[0], "nombre aproximado"
    if len(partial) > 1:
        return None, f"{len(partial)} coincidencias; usa un nombre más específico"
    return None, "no encontrado"


def resolver_foto_poste(
    zona: str,
    poste: str,
    tipo: str,
    indexed_files: List[dict],
    explicit_ref: str = "",
) -> Tuple[Optional[dict], str]:
    """
    Resuelve una evidencia por ZONA + N° POSTE + sufijo A/D.

    Prioridad segura:
    1) Enlace/ID directo de Drive indicado explícitamente (override intencional).
    2) Búsqueda automática dentro de la carpeta de la ZONA: POSTE_A / POSTE_D.
    3) Coincidencia global única por nombre, si no existe en la carpeta de zona.
    4) Referencia de texto de las columnas FOTO ANTES/FOTO DESPUES como último recurso.

    Esto evita que un código genérico como ``116_A`` tome por error una foto homónima
    ubicada en otra carpeta de Drive.
    """
    explicit_clean = "" if explicit_ref is None else str(explicit_ref).strip()

    # Si el usuario pegó un enlace o un fileId de Drive, sí se considera un override
    # intencional y tiene prioridad sobre la convención automática.
    if explicit_clean and explicit_clean.lower() != "nan":
        direct_id = extract_drive_file_id(explicit_clean)
        if direct_id:
            meta, err = get_drive_metadata(direct_id)
            if meta:
                return meta, "referencia explícita por enlace/ID"

    zona_key = normalize_text_token(zona)
    poste_code = normalize_poste_code(poste)
    sufijo = "A" if normalize_text_token(tipo).startswith("A") else "D"
    target_key = normalize_photo_key(f"{poste_code}_{sufijo}")
    if not target_key:
        return None, "poste inválido"

    exact = [
        f for f in indexed_files
        if str(f.get("mimeType", "")).startswith(IMAGE_MIME_PREFIX)
        and f.get("key_stem") == target_key
    ]

    # La carpeta de ZONA manda. Esto es lo que hace segura la automatización.
    if zona_key:
        exact_zone = [f for f in exact if zona_key in f.get("folder_keys", [])]
        if exact_zone:
            exact_zone.sort(key=lambda x: str(x.get("modifiedTime", "")), reverse=True)
            if len(exact_zone) == 1:
                return exact_zone[0], "zona + nombre exacto"
            return exact_zone[0], f"zona + nombre exacto; {len(exact_zone)} duplicados, se usa el más reciente"

    # Si no se encontró en la carpeta de zona, solo aceptamos una coincidencia global
    # cuando sea inequívoca.
    if len(exact) == 1:
        return exact[0], "nombre exacto único (fuera de carpeta de zona esperada)"
    if len(exact) > 1:
        # Antes de rendirnos, una referencia de texto puede ser más específica.
        if explicit_clean and explicit_clean.lower() != "nan":
            meta, reason = resolver_foto(explicit_clean, indexed_files)
            if meta:
                return meta, f"referencia explícita de respaldo ({reason})"
        return None, f"{len(exact)} archivos con el mismo nombre; revisar carpeta de zona"

    # Compatibilidad con la hoja anterior: FOTO ANTES / FOTO DESPUES se usan solo
    # como respaldo cuando la nomenclatura automática no encontró la evidencia.
    if explicit_clean and explicit_clean.lower() != "nan":
        meta, reason = resolver_foto(explicit_clean, indexed_files)
        if meta:
            return meta, f"referencia explícita de respaldo ({reason})"

    return None, f"no se encontró {poste_code}_{sufijo} en la zona {zona}"


def drive_preview_url(file_id: str) -> str:
    return f"https://drive.google.com/file/d/{file_id}/preview"


def drive_view_url(file_id: str) -> str:
    return f"https://drive.google.com/file/d/{file_id}/view"


def render_photo_panel(label: str, raw_value: str, indexed_files: List[dict]):
    st.markdown(f"#### 📸 {label}")
    raw_value = "" if raw_value is None else str(raw_value).strip()
    if not raw_value or raw_value.lower() == "nan":
        st.info("No hay código o enlace registrado.")
        return

    st.caption(f"Referencia en Google Sheets: {raw_value}")
    meta, metodo = resolver_foto(raw_value, indexed_files)

    if meta:
        file_id = meta.get("id")
        mime = meta.get("mimeType", "")
        file_name = meta.get("name", raw_value)

        if mime.startswith(IMAGE_MIME_PREFIX) and has_service_account_secrets() and GOOGLE_DRIVE_LIBS_OK:
            data, mime_download, err_or_name = descargar_archivo_drive(file_id)
            if data:
                st.image(data, caption=file_name, use_container_width=True)
            else:
                st.warning(f"Se encontró el archivo, pero no se pudo descargar: {err_or_name}")
                st.markdown(
                    f'<iframe src="{drive_preview_url(file_id)}" width="100%" height="420" '
                    'style="border:1px solid #ced4da;border-radius:8px;" allow="autoplay"></iframe>',
                    unsafe_allow_html=True,
                )
        else:
            # Fallback útil para archivos públicos o cuando no se configuró Service Account.
            st.markdown(
                f'<iframe src="{drive_preview_url(file_id)}" width="100%" height="420" '
                'style="border:1px solid #ced4da;border-radius:8px;" allow="autoplay"></iframe>',
                unsafe_allow_html=True,
            )

        st.markdown(f"[Abrir archivo en Google Drive]({drive_view_url(file_id)})")
        st.caption(f"Resuelto por: {metodo} · Archivo: {file_name}")
        return

    # Si el valor era enlace/ID y no hay credenciales, todavía intenta el preview del navegador.
    file_id = extract_drive_file_id(raw_value)
    if file_id:
        st.markdown(
            f'<iframe src="{drive_preview_url(file_id)}" width="100%" height="420" '
            'style="border:1px solid #ced4da;border-radius:8px;" allow="autoplay"></iframe>',
            unsafe_allow_html=True,
        )
        st.caption("Vista directa por enlace/ID. Para archivos privados configura la Service Account.")
    else:
        st.warning(f"Foto no localizada automáticamente ({metodo}).")


def render_photo_automatica(
    label: str,
    zona: str,
    poste: str,
    tipo: str,
    indexed_files: List[dict],
    explicit_ref: str = "",
):
    st.markdown(f"#### 📸 {label}")
    meta, metodo = resolver_foto_poste(zona, poste, tipo, indexed_files, explicit_ref=explicit_ref)
    if not meta:
        esperado = f"{normalize_poste_code(poste)}_{'A' if tipo.upper().startswith('A') else 'D'}"
        st.warning(f"No encontrada. Esperado en `{zona}`: `{esperado}.jpg/.jpeg/.png`")
        st.caption(metodo)
        return

    file_id = meta.get("id")
    file_name = meta.get("name", "")
    mime = meta.get("mimeType", "")
    carpeta = meta.get("relative_folder", "")
    st.caption(f"{carpeta}/{file_name} · {metodo}" if carpeta else f"{file_name} · {metodo}")

    if mime.startswith(IMAGE_MIME_PREFIX) and has_service_account_secrets() and GOOGLE_DRIVE_LIBS_OK:
        data, _, err_or_name = descargar_archivo_drive(file_id)
        if data:
            st.image(data, caption=file_name, use_container_width=True)
        else:
            st.warning(f"Se encontró el archivo, pero no se pudo descargar: {err_or_name}")
    else:
        st.markdown(
            f'<iframe src="{drive_preview_url(file_id)}" width="100%" height="420" '
            'style="border:1px solid #ced4da;border-radius:8px;" allow="autoplay"></iframe>',
            unsafe_allow_html=True,
        )
    st.markdown(f"[Abrir archivo en Google Drive]({drive_view_url(file_id)})")


# -----------------------------------------------------------------------------
# CALIDAD Y UTILIDADES DE DATOS
# -----------------------------------------------------------------------------
def resaltar_filas(row):
    estado = str(row.get("ESTADO", "")).upper()
    if estado == "CONFORME":
        return ["background-color:#d9ead3;color:#274e13;font-weight:bold"] * len(row)
    if estado == "ATENDIDO":
        return ["background-color:#fff2cc;color:#7f6000;font-weight:bold"] * len(row)
    if estado == "PENDIENTE":
        return ["background-color:#fce8e6;color:#8a1c13"] * len(row)
    return [""] * len(row)


def safe_text(value) -> str:
    if pd.isna(value):
        return ""
    return str(value)


def coerce_coordinates(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col in ["LATITUD", "LONGITUD"]:
        if col in out.columns:
            out[col] = (
                out[col]
                .astype(str)
                .str.replace(",", ".", regex=False)
            )
            out[col] = pd.to_numeric(out[col], errors="coerce")
    return out


def detectar_columnas_coordenadas(df: pd.DataFrame) -> Tuple[Optional[str], Optional[str]]:
    cols_upper = {str(c).upper().strip(): c for c in df.columns}
    lat_candidates = ["LATITUD", "LAT", "Y"]
    lon_candidates = ["LONGITUD", "LON", "LONG", "X"]
    lat = next((cols_upper[x] for x in lat_candidates if x in cols_upper), None)
    lon = next((cols_upper[x] for x in lon_candidates if x in cols_upper), None)
    return lat, lon


# -----------------------------------------------------------------------------
# EXPORTACIÓN EXCEL
# -----------------------------------------------------------------------------
def generar_excel_estilizado(data_df: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Control de Postes"

    headers = list(data_df.columns)
    ws.append(headers)

    header_fill = PatternFill(start_color="343A40", end_color="343A40", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="DEE2E6"),
        right=Side(style="thin", color="DEE2E6"),
        top=Side(style="thin", color="DEE2E6"),
        bottom=Side(style="thin", color="DEE2E6"),
    )

    fills = {
        "CONFORME": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),
        "ATENDIDO": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        "PENDIENTE": PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid"),
    }

    for col_num, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border

    estado_idx = headers.index("ESTADO") if "ESTADO" in headers else None

    for row_idx, row in enumerate(data_df.itertuples(index=False, name=None), start=2):
        ws.append([safe_text(v) for v in row])
        estado_val = str(row[estado_idx]).upper() if estado_idx is not None else ""
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.border = thin_border
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if estado_val in fills:
                cell.fill = fills[estado_val]

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for i, col in enumerate(ws.columns, start=1):
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(max_len + 3, 12), 45)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# -----------------------------------------------------------------------------
# EXPORTACIÓN PDF
# -----------------------------------------------------------------------------
def _matplotlib_summary_images(data_df: pd.DataFrame):
    img_buf_1, img_buf_2 = None, None

    if "ESTADO" in data_df.columns and len(data_df) > 0:
        fig, ax = plt.subplots(figsize=(4.2, 2.2))
        conteo = data_df["ESTADO"].value_counts()
        colores = [COLOR_ESTADOS.get(x, "#6c757d") for x in conteo.index]
        wedges, _, _ = ax.pie(
            conteo,
            labels=None,
            colors=colores,
            autopct="%1.1f%%",
            startangle=90,
            pctdistance=0.55,
            textprops={"fontsize": 8, "weight": "bold", "color": "white"},
        )
        ax.legend(
            wedges,
            conteo.index,
            title="Estado",
            loc="center left",
            bbox_to_anchor=(0.95, 0.5),
            fontsize=8,
            title_fontsize=8,
        )
        ax.set_title("Distribución por Estado", fontsize=9, fontweight="bold")
        plt.tight_layout()
        img_buf_1 = io.BytesIO()
        plt.savefig(img_buf_1, format="png", dpi=150, bbox_inches="tight")
        img_buf_1.seek(0)
        plt.close(fig)

    if "ZONA" in data_df.columns and "ESTADO" in data_df.columns and len(data_df) > 0:
        fig, ax = plt.subplots(figsize=(4.8, 2.2))
        df_zona_est = data_df.groupby(["ZONA", "ESTADO"]).size().unstack(fill_value=0)
        for est in ESTADOS_BASE:
            if est not in df_zona_est.columns:
                df_zona_est[est] = 0
        df_zona_est = df_zona_est[ESTADOS_BASE]

        bottom = None
        for estado in ESTADOS_BASE:
            values = df_zona_est[estado].values
            ax.bar(
                df_zona_est.index.astype(str),
                values,
                bottom=bottom,
                label=estado,
                color=COLOR_ESTADOS[estado],
                width=0.45,
            )
            bottom = values if bottom is None else bottom + values

        ax.legend(title="Estado", loc="upper right", fontsize=7, title_fontsize=7)
        ax.set_title("Observados por Zonas", fontsize=9, fontweight="bold")
        ax.tick_params(axis="x", rotation=10, labelsize=8)
        ax.tick_params(axis="y", labelsize=8)
        plt.tight_layout()
        img_buf_2 = io.BytesIO()
        plt.savefig(img_buf_2, format="png", dpi=150, bbox_inches="tight")
        img_buf_2.seek(0)
        plt.close(fig)

    return img_buf_1, img_buf_2


def _prepare_pdf_photo(file_ref: str, indexed_files: List[dict], max_width=250, max_height=170):
    meta, _ = resolver_foto(file_ref, indexed_files)
    return _prepare_pdf_photo_meta(meta, max_width=max_width, max_height=max_height)


def _prepare_pdf_photo_meta(meta: Optional[dict], max_width=250, max_height=170):
    if not meta or not str(meta.get("mimeType", "")).startswith(IMAGE_MIME_PREFIX):
        return None
    data, _, _ = descargar_archivo_drive(meta["id"])
    if not data:
        return None
    try:
        image = RLImage(io.BytesIO(data))
        image._restrictSize(max_width, max_height)
        return image
    except Exception:
        return None


def generar_pdf_ejecutivo(
    data_df: pd.DataFrame,
    total: int,
    pend: int,
    aten: int,
    conf: int,
    indexed_files: Optional[List[dict]] = None,
    incluir_anexo_fotos: bool = False,
    max_fotos: int = 20,
) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=26,
        leftMargin=26,
        topMargin=25,
        bottomMargin=25,
    )
    elements = []

    styles = getSampleStyleSheet()
    titulo_estilo = ParagraphStyle(
        "Titulo",
        parent=styles["Heading1"],
        fontSize=15,
        alignment=1,
        textColor=colors.HexColor("#212529"),
    )
    sub_estilo = ParagraphStyle(
        "Sub",
        parent=styles["Normal"],
        fontSize=9,
        alignment=1,
        textColor=colors.HexColor("#6c757d"),
    )
    seccion_estilo = ParagraphStyle(
        "Sec",
        parent=styles["Heading2"],
        fontSize=11,
        textColor=colors.HexColor("#343a40"),
        spaceBefore=8,
        spaceAfter=5,
    )
    celda = ParagraphStyle("Celda", parent=styles["Normal"], fontSize=6.7, leading=8)
    cab = ParagraphStyle(
        "Cabecera",
        parent=styles["Normal"],
        fontSize=7,
        fontName="Helvetica-Bold",
        textColor=colors.whitesmoke,
        alignment=1,
    )

    elements.append(Paragraph("<b>REPORTE EJECUTIVO - CONTROL DE POSTES OBSERVADOS</b>", titulo_estilo))
    elements.append(Paragraph("Monitoreo de avance y levantamiento de observaciones en obra", sub_estilo))
    elements.append(Spacer(1, 8))

    avance = ((aten + conf) / total * 100) if total else 0
    kpi_text = (
        f"<b>Total:</b> {total} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<font color='#d9534f'><b>Pendientes:</b> {pend}</font> &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<font color='#7f6000'><b>Atendidos:</b> {aten}</font> &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<font color='#274e13'><b>Conformes:</b> {conf}</font> &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>Levantamiento:</b> {avance:.1f}%"
    )
    elements.append(Paragraph(kpi_text, ParagraphStyle("KPI", parent=styles["Normal"], fontSize=9, alignment=1)))
    elements.append(Spacer(1, 10))

    img1, img2 = _matplotlib_summary_images(data_df)
    graph_cells = []
    if img1:
        graph_cells.append(RLImage(img1, width=220, height=120))
    if img2:
        graph_cells.append(RLImage(img2, width=250, height=120))
    if graph_cells:
        while len(graph_cells) < 2:
            graph_cells.append("")
        tg = Table([graph_cells], colWidths=[360, 360])
        tg.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER")]))
        elements.append(tg)

    elements.append(Spacer(1, 8))
    elements.append(Paragraph("<b>Detalle de registros</b>", seccion_estilo))

    # Para el PDF, evita que columnas de fotos/links muy extensas vuelvan ilegible la tabla.
    preferred = [
        "ZONA",
        "N° POSTE",
        "TIPO / ALTURA",
        "TERRENO",
        "JUSTIFICACIÓN",
        "OBSERVACIÓN / ACCIÓN",
        "ESTADO",
    ]
    pdf_cols = [c for c in preferred if c in data_df.columns]
    if not pdf_cols:
        pdf_cols = list(data_df.columns[:7])

    table_data = [[Paragraph(str(c), cab) for c in pdf_cols]]
    for _, row in data_df[pdf_cols].iterrows():
        table_data.append([Paragraph(safe_text(v), celda) for v in row])

    # Anchos pensados para landscape letter.
    width_map = {
        "ZONA": 70,
        "N° POSTE": 55,
        "TIPO / ALTURA": 70,
        "TERRENO": 70,
        "JUSTIFICACIÓN": 145,
        "OBSERVACIÓN / ACCIÓN": 220,
        "ESTADO": 65,
    }
    col_widths = [width_map.get(c, 90) for c in pdf_cols]

    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#343a40")),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#dee2e6")),
            ]
        )
    )

    # Color por estado fila a fila.
    if "ESTADO" in pdf_cols:
        est_col = pdf_cols.index("ESTADO")
        for r_idx, (_, row) in enumerate(data_df[pdf_cols].iterrows(), start=1):
            est = str(row["ESTADO"]).upper()
            bg = {
                "CONFORME": "#d9ead3",
                "ATENDIDO": "#fff2cc",
                "PENDIENTE": "#fce8e6",
            }.get(est)
            if bg:
                t.setStyle(TableStyle([("BACKGROUND", (0, r_idx), (-1, r_idx), colors.HexColor(bg))]))

    elements.append(t)

    # Anexo opcional automático por ZONA + N° POSTE.
    if incluir_anexo_fotos and indexed_files and "N° POSTE" in data_df.columns:
        candidatos = []
        for _, row in data_df.iterrows():
            zona = safe_text(row.get("ZONA", ""))
            poste = normalize_poste_code(row.get("N° POSTE", ""))
            ref_a = safe_text(row.get("FOTO ANTES", "")) if "FOTO ANTES" in data_df.columns else ""
            ref_d = safe_text(row.get("FOTO DESPUES", "")) if "FOTO DESPUES" in data_df.columns else ""
            meta_a, _ = resolver_foto_poste(zona, poste, "ANTES", indexed_files, explicit_ref=ref_a)
            meta_d, _ = resolver_foto_poste(zona, poste, "DESPUES", indexed_files, explicit_ref=ref_d)
            if meta_a or meta_d:
                candidatos.append((zona, poste, meta_a, meta_d))
            if len(candidatos) >= max_fotos:
                break

        if candidatos:
            elements.append(PageBreak())
            elements.append(Paragraph("<b>Anexo fotográfico</b>", seccion_estilo))
            for zona, poste, meta_a, meta_d in candidatos:
                elements.append(Paragraph(f"<b>Poste {poste}</b> - {zona}", styles["Normal"]))
                photos = [
                    _prepare_pdf_photo_meta(meta_a) if meta_a else "",
                    _prepare_pdf_photo_meta(meta_d) if meta_d else "",
                ]
                pt = Table([photos, ["ANTES", "DESPUÉS"]], colWidths=[330, 330])
                pt.setStyle(TableStyle([
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 1), (-1, 1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ]))
                elements.append(pt)
                elements.append(Spacer(1, 6))

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()


# -----------------------------------------------------------------------------
# CABECERA
# -----------------------------------------------------------------------------
# Se deja un margen superior propio para evitar que la barra flotante de
# Streamlit recorte visualmente la parte superior del logo en pantallas anchas.
st.markdown("<div style='height:14px;'></div>", unsafe_allow_html=True)

col_logo, col_title = st.columns([1, 4])
with col_logo:
    logo_path = Path("quantum.png")
    if logo_path.exists():
        st.markdown("<div style='height:4px;'></div>", unsafe_allow_html=True)
        st.image(str(logo_path), width=265)
    else:
        st.markdown("### 🏗️")
        st.caption("Logo quantum.png no encontrado")

with col_title:
    st.markdown(
        """
        <div style='display:flex;flex-direction:column;justify-content:center;height:100%;text-align:center;padding-top:8px;'>
            <h2 style='color:#212529;margin-bottom:0;font-size:calc(1.3rem + 1vw);'>📊 DASHBOARD CONTROL DE POSTES OBSERVADOS</h2>
            <p style='color:#6c757d;margin-top:5px;font-size:calc(0.9rem + 0.3vw);'>Monitoreo en tiempo real, trazabilidad fotográfica y levantamiento de observaciones.</p>
        </div>
        """,
        unsafe_allow_html=True,
    )

st.markdown("---")

# -----------------------------------------------------------------------------
# CARGA DE DATOS
# -----------------------------------------------------------------------------
sheet_url = get_sheet_url()
df_raw, error_detallado = cargar_datos_gsheets(sheet_url)

if df_raw is None or len(df_raw) == 0:
    st.error("No se pudo cargar la información desde Google Sheets.")
    if error_detallado:
        st.info(f"Detalle técnico: {error_detallado}")
    st.stop()

try:
    df = preparar_dataframe(df_raw)
except Exception as e:
    st.error(f"No se pudo preparar la estructura de datos: {e}")
    st.stop()

# -----------------------------------------------------------------------------
# ÍNDICE DE DRIVE
# -----------------------------------------------------------------------------
folder_cfg = get_drive_folder_config()
drive_files_antes, drive_error_antes = ([], None)
drive_files_despues, drive_error_despues = ([], None)

if has_service_account_secrets() and GOOGLE_DRIVE_LIBS_OK:
    if folder_cfg["antes"]:
        drive_files_antes, drive_error_antes = indexar_drive(tuple(folder_cfg["antes"]))
    if folder_cfg["despues"] == folder_cfg["antes"]:
        drive_files_despues, drive_error_despues = drive_files_antes, drive_error_antes
    elif folder_cfg["despues"]:
        drive_files_despues, drive_error_despues = indexar_drive(tuple(folder_cfg["despues"]))

# -----------------------------------------------------------------------------
# SIDEBAR
# -----------------------------------------------------------------------------
st.sidebar.header("🔍 Filtros de búsqueda")

zona_sel = "TODOS"
if "ZONA" in df.columns:
    zonas = ["TODOS"] + sorted(df["ZONA"].dropna().astype(str).unique().tolist())
    zona_sel = st.sidebar.selectbox("Zona / Sector", zonas)

estado_sel = "TODOS"
if "ESTADO" in df.columns:
    estados = ["TODOS"] + sorted(df["ESTADO"].dropna().astype(str).unique().tolist())
    estado_sel = st.sidebar.selectbox("Estado", estados)

terreno_sel = "TODOS"
if "TERRENO" in df.columns:
    terrenos = ["TODOS"] + sorted(df["TERRENO"].dropna().astype(str).unique().tolist())
    terreno_sel = st.sidebar.selectbox("Tipo de terreno", terrenos)

texto_busqueda = st.sidebar.text_input("Buscar poste / observación", placeholder="Ej.: 116 o cimentación")

st.sidebar.markdown("---")
st.sidebar.subheader("☁️ Estado Google Drive")
if not GOOGLE_DRIVE_LIBS_OK:
    st.sidebar.error("Faltan librerías de Google Drive")
elif not has_service_account_secrets():
    st.sidebar.warning("Service Account no configurada")
elif not folder_cfg["root"] and not folder_cfg["antes"] and not folder_cfg["despues"]:
    st.sidebar.warning("Faltan IDs de carpetas")
elif drive_error_antes or drive_error_despues:
    st.sidebar.error("Error al indexar Drive")
    with st.sidebar.expander("Ver detalle"):
        st.write(drive_error_antes or drive_error_despues)
else:
    st.sidebar.success("Drive conectado")
    if folder_cfg["despues"] == folder_cfg["antes"]:
        st.sidebar.caption(f"Archivos de Drive indexados: {len(drive_files_antes)}")
    else:
        st.sidebar.caption(
            f"Archivos indexados: antes {len(drive_files_antes)} · después {len(drive_files_despues)}"
        )

if st.sidebar.button("🔄 Refrescar datos e índice", use_container_width=True):
    st.cache_data.clear()
    st.rerun()

# -----------------------------------------------------------------------------
# FILTROS
# -----------------------------------------------------------------------------
df_filtrado = df.copy()
if zona_sel != "TODOS" and "ZONA" in df.columns:
    df_filtrado = df_filtrado[df_filtrado["ZONA"].astype(str) == zona_sel]
if estado_sel != "TODOS" and "ESTADO" in df.columns:
    df_filtrado = df_filtrado[df_filtrado["ESTADO"].astype(str) == estado_sel]
if terreno_sel != "TODOS" and "TERRENO" in df.columns:
    df_filtrado = df_filtrado[df_filtrado["TERRENO"].astype(str) == terreno_sel]
if texto_busqueda.strip():
    q = texto_busqueda.strip().lower()
    mask = df_filtrado.astype(str).apply(lambda col: col.str.lower().str.contains(q, regex=False)).any(axis=1)
    df_filtrado = df_filtrado[mask]

# -----------------------------------------------------------------------------
# KPIs
# -----------------------------------------------------------------------------
total_postes = len(df_filtrado)
pendientes = len(df_filtrado[df_filtrado["ESTADO"] == "PENDIENTE"]) if "ESTADO" in df_filtrado.columns else 0
atendidos = len(df_filtrado[df_filtrado["ESTADO"] == "ATENDIDO"]) if "ESTADO" in df_filtrado.columns else 0
conformes = len(df_filtrado[df_filtrado["ESTADO"] == "CONFORME"]) if "ESTADO" in df_filtrado.columns else 0
avance = ((atendidos + conformes) / total_postes * 100) if total_postes else 0
conformidad = (conformes / total_postes * 100) if total_postes else 0

sin_foto_despues = 0
if has_service_account_secrets() and drive_files_despues and "N° POSTE" in df_filtrado.columns:
    for _, row in df_filtrado.iterrows():
        zona = safe_text(row.get("ZONA", ""))
        poste = normalize_poste_code(row.get("N° POSTE", ""))
        ref = safe_text(row.get("FOTO DESPUES", "")) if "FOTO DESPUES" in df_filtrado.columns else ""
        meta, _ = resolver_foto_poste(zona, poste, "DESPUES", drive_files_despues, explicit_ref=ref)
        if not meta:
            sin_foto_despues += 1
else:
    if "FOTO DESPUES" in df_filtrado.columns:
        vals = df_filtrado["FOTO DESPUES"].astype(str).str.strip()
        sin_foto_despues = int(((vals == "") | (vals.str.lower() == "nan")).sum())

kcols = st.columns(6)
kpis = [
    ("Observados", f"🔍 {total_postes}", "Registros filtrados", "#212529"),
    ("Pendientes", f"⚠️ {pendientes}", "Por atender", "#d9534f"),
    ("Atendidos", f"🔧 {atendidos}", "Levantamiento ejecutado", "#b07900"),
    ("Conformes", f"🏆 {conformes}", "Aceptados", "#22863a"),
    ("Levantamiento", f"📈 {avance:.1f}%", "Atendidos + conformes", "#0d6efd"),
    ("Sin foto después", f"📷 {sin_foto_despues}", "Pendiente evidencia", "#6c757d"),
]
for col, (title, value, sub, color) in zip(kcols, kpis):
    with col:
        st.markdown(
            f"""
            <div class='metric-card'>
                <div class='metric-title'>{title}</div>
                <div class='metric-value' style='color:{color};'>{value}</div>
                <div class='metric-sub'>{sub}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

st.markdown("<br>", unsafe_allow_html=True)

# -----------------------------------------------------------------------------
# GRÁFICOS
# -----------------------------------------------------------------------------
col_g1, col_g2 = st.columns(2)

with col_g1:
    if "ZONA" in df_filtrado.columns and "ESTADO" in df_filtrado.columns and len(df_filtrado):
        st.markdown('<div class="centered-subheader">🗺️ Observados por Zonas</div>', unsafe_allow_html=True)
        dz = df_filtrado.groupby(["ZONA", "ESTADO"]).size().unstack(fill_value=0).reset_index()
        fig = go.Figure()
        for estado in ESTADOS_BASE:
            if estado in dz.columns:
                fig.add_trace(
                    go.Bar(
                        name=estado,
                        x=dz["ZONA"],
                        y=dz[estado],
                        text=dz[estado],
                        textposition="inside",
                        textfont=dict(color="white", size=12),
                        marker_color=COLOR_ESTADOS[estado],
                        marker_line=dict(color="#111111", width=1),
                    )
                )
        fig.update_layout(
            barmode="stack",
            margin=dict(t=20, b=0, l=0, r=0),
            height=330,
            paper_bgcolor="rgba(0,0,0,0)",
            plot_bgcolor="rgba(0,0,0,0)",
            xaxis=dict(title="Zona", showgrid=False),
            yaxis=dict(title="Cantidad", gridcolor="#e6e6e6"),
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="center", x=0.5),
        )
        st.plotly_chart(fig, use_container_width=True, config={"displayModeBar": False})

with col_g2:
    if "ESTADO" in df_filtrado.columns and len(df_filtrado):
        st.markdown('<div class="centered-subheader">📌 Distribución del Estado</div>', unsafe_allow_html=True)
        ce = df_filtrado["ESTADO"].value_counts().reset_index()
        ce.columns = ["ESTADO", "CANTIDAD"]
        fig = px.pie(
            ce,
            names="ESTADO",
            values="CANTIDAD",
            hole=0.4,
            color="ESTADO",
            color_discrete_map=COLOR_ESTADOS,
        )
        fig.update_traces(textposition="inside", textinfo="percent+label", marker=dict(line=dict(color="#000", width=1)))
        fig.update_layout(
            margin=dict(t=0, b=0, l=0, r=0),
            height=330,
            paper_bgcolor="rgba(0,0,0,0)",
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="center", x=0.5),
        )
        st.plotly_chart(fig, use_container_width=True, config={"displayModeBar": False})

if "TERRENO" in df_filtrado.columns and "ESTADO" in df_filtrado.columns and len(df_filtrado):
    st.markdown('<div class="centered-subheader">🌍 Clasificación por Tipo de Terreno</div>', unsafe_allow_html=True)
    dt = df_filtrado.groupby(["TERRENO", "ESTADO"]).size().unstack(fill_value=0).reset_index()
    fig = go.Figure()
    for estado in ESTADOS_BASE:
        if estado in dt.columns:
            fig.add_trace(
                go.Bar(
                    name=estado,
                    x=dt["TERRENO"],
                    y=dt[estado],
                    text=dt[estado],
                    textposition="inside",
                    textfont=dict(color="white", size=12),
                    marker_color=COLOR_ESTADOS[estado],
                    marker_line=dict(color="#111", width=1),
                )
            )
    fig.update_layout(
        barmode="stack",
        height=330,
        margin=dict(t=20, b=0, l=0, r=0),
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        xaxis=dict(title="Tipo de terreno", showgrid=False),
        yaxis=dict(title="Cantidad", gridcolor="#e6e6e6"),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="center", x=0.5),
    )
    st.plotly_chart(fig, use_container_width=True, config={"displayModeBar": False})

# -----------------------------------------------------------------------------
# MAPA OPCIONAL
# -----------------------------------------------------------------------------
lat_col, lon_col = detectar_columnas_coordenadas(df_filtrado)
if lat_col and lon_col:
    map_df = df_filtrado.copy()
    map_df[lat_col] = pd.to_numeric(map_df[lat_col].astype(str).str.replace(",", ".", regex=False), errors="coerce")
    map_df[lon_col] = pd.to_numeric(map_df[lon_col].astype(str).str.replace(",", ".", regex=False), errors="coerce")
    map_df = map_df.dropna(subset=[lat_col, lon_col])
    if len(map_df):
        st.markdown('<div class="centered-subheader">📍 Ubicación de Postes</div>', unsafe_allow_html=True)
        st.caption(f"{len(map_df)} poste(s) con coordenadas válidas en el filtro actual.")
        hover = [c for c in ["N° POSTE", "ZONA", "ESTADO", "TERRENO", "OBSERVACIÓN / ACCIÓN"] if c in map_df.columns]
        fig_map = px.scatter_map(
            map_df,
            lat=lat_col,
            lon=lon_col,
            color="ESTADO" if "ESTADO" in map_df.columns else None,
            color_discrete_map=COLOR_ESTADOS,
            hover_data=hover,
            zoom=12,
            height=480,
        )
        fig_map.update_layout(map_style="open-street-map", margin=dict(l=0, r=0, t=0, b=0))
        st.plotly_chart(fig_map, use_container_width=True, config={"displayModeBar": False})
else:
    with st.expander("📍 ¿Quieres activar el mapa de postes?"):
        st.info(
            "Agrega columnas LATITUD y LONGITUD en Google Sheets. El dashboard las detectará automáticamente "
            "y mostrará únicamente los postes que tengan coordenadas válidas."
        )

st.markdown("---")

# -----------------------------------------------------------------------------
# CONSULTA INDIVIDUAL + FOTOS AUTOMÁTICAS POR ZONA / POSTE
# -----------------------------------------------------------------------------
st.subheader("🔎 Consulta individual de poste y fotografías")

if "N° POSTE" in df.columns:
    opciones = []
    for idx, row in df.iterrows():
        zona = safe_text(row.get("ZONA", ""))
        poste = normalize_poste_code(row.get("N° POSTE", ""))
        opciones.append((f"{zona} · Poste {poste}", idx))

    labels = ["-- Seleccionar --"] + [x[0] for x in opciones]
    seleccion = st.selectbox("Seleccione o busque zona y número de poste", labels)

    if seleccion != "-- Seleccionar --":
        selected_pos = labels.index(seleccion) - 1
        selected_index = opciones[selected_pos][1]
        datos_poste = df.loc[[selected_index]]
        row = datos_poste.iloc[0]
        zona = safe_text(row.get("ZONA", ""))
        poste = normalize_poste_code(row.get("N° POSTE", ""))

        st.dataframe(datos_poste.style.apply(resaltar_filas, axis=1), use_container_width=True)

        # Mensaje persistente después del rerun provocado por una carga exitosa.
        flash_key = f"upload_flash_{selected_index}"
        if flash_key in st.session_state:
            st.success(st.session_state.pop(flash_key))

        with st.expander("📤 Cargar o reemplazar fotografías de este poste"):
            upload_pin_cfg = get_upload_pin()

            if not has_service_account_secrets() or not GOOGLE_DRIVE_LIBS_OK:
                st.warning("Google Drive no está configurado para realizar cargas.")
            elif not folder_cfg["root"]:
                st.warning("Falta configurar [drive].root_folder_ids en Streamlit Secrets.")
            elif not upload_pin_cfg:
                st.warning(
                    "Por seguridad, la carga está deshabilitada hasta configurar un PIN. "
                    "Agrega en Streamlit Secrets, dentro de [app], una línea como: "
                    'upload_pin = "TU_CLAVE_DE_CARGA"'
                )
            else:
                pin_ingresado = st.text_input(
                    "Clave de carga",
                    type="password",
                    key=f"upload_pin_{selected_index}",
                    help="Esta clave evita que cualquier visitante de una app pública pueda modificar las evidencias.",
                )
                autorizado = hmac.compare_digest(pin_ingresado, upload_pin_cfg) if pin_ingresado else False

                if pin_ingresado and not autorizado:
                    st.error("Clave de carga incorrecta.")

                if autorizado:
                    st.caption(
                        f"Destino automático: {zona}/ · Poste {poste}. "
                        "ANTES se guarda como _A y DESPUÉS como _D. Si ya existe, se reemplaza."
                    )
                    with st.form(f"form_subir_fotos_{selected_index}", clear_on_submit=True):
                        up1, up2 = st.columns(2)
                        with up1:
                            foto_antes_upload = st.file_uploader(
                                "📷 Foto ANTES",
                                type=["jpg", "jpeg", "png", "webp"],
                                key=f"foto_antes_upload_{selected_index}",
                            )
                        with up2:
                            foto_despues_upload = st.file_uploader(
                                "📷 Foto DESPUÉS",
                                type=["jpg", "jpeg", "png", "webp"],
                                key=f"foto_despues_upload_{selected_index}",
                            )

                        enviar_fotos = st.form_submit_button(
                            "☁️ Guardar fotografías en Google Drive",
                            use_container_width=True,
                        )

                    if enviar_fotos:
                        if foto_antes_upload is None and foto_despues_upload is None:
                            st.warning("Selecciona al menos una fotografía.")
                        else:
                            resultados_carga = []
                            errores_carga = []
                            with st.spinner("Guardando fotografías en Google Drive..."):
                                for tipo_carga, archivo_carga in [
                                    ("ANTES", foto_antes_upload),
                                    ("DESPUES", foto_despues_upload),
                                ]:
                                    if archivo_carga is None:
                                        continue
                                    meta_up, accion_up, error_up = subir_o_reemplazar_foto_poste(
                                        zona,
                                        poste,
                                        tipo_carga,
                                        archivo_carga,
                                        folder_cfg["root"],
                                    )
                                    if error_up:
                                        errores_carga.append(f"{tipo_carga}: {error_up}")
                                    else:
                                        resultados_carga.append(
                                            f"{tipo_carga}: {meta_up.get('name', '')} {accion_up}"
                                        )

                            if errores_carga:
                                for msg_error in errores_carga:
                                    st.error(msg_error)
                            if resultados_carga:
                                limpiar_cache_drive()
                                st.session_state[flash_key] = "Fotografías guardadas correctamente. " + " · ".join(resultados_carga)
                                st.rerun()

        # Encabezado técnico del registro fotográfico oculto en la interfaz.
        # Las fotografías ANTES / DESPUÉS siguen mostrándose normalmente.
        ref_a = safe_text(row.get("FOTO ANTES", "")) if "FOTO ANTES" in df.columns else ""
        ref_d = safe_text(row.get("FOTO DESPUES", "")) if "FOTO DESPUES" in df.columns else ""

        c1, c2 = st.columns(2)
        with c1:
            render_photo_automatica("ANTES", zona, poste, "ANTES", drive_files_antes, explicit_ref=ref_a)
        with c2:
            render_photo_automatica("DESPUÉS", zona, poste, "DESPUES", drive_files_despues, explicit_ref=ref_d)
else:
    st.info("No se encontró la columna N° POSTE en la fuente de datos.")

st.markdown("---")

# -----------------------------------------------------------------------------
# CONTROL DE CALIDAD DE FOTOS
# -----------------------------------------------------------------------------
with st.expander("🧪 Control de calidad de evidencias fotográficas"):
    if "N° POSTE" not in df_filtrado.columns:
        st.info("No se encontró la columna N° POSTE para validar evidencias.")
    else:
        resultados = []
        for _, row in df_filtrado.iterrows():
            poste = normalize_poste_code(row.get("N° POSTE", ""))
            zona = safe_text(row.get("ZONA", ""))
            for tipo, col, idx_files in [
                ("ANTES", "FOTO ANTES", drive_files_antes),
                ("DESPUÉS", "FOTO DESPUES", drive_files_despues),
            ]:
                ref = safe_text(row.get(col, "")) if col in df_filtrado.columns else ""
                if has_service_account_secrets() and idx_files:
                    meta, reason = resolver_foto_poste(zona, poste, tipo, idx_files, explicit_ref=ref)
                    status = "OK" if meta else f"NO ENCONTRADA ({reason})"
                    archivo = meta.get("name", "") if meta else ""
                    carpeta = meta.get("relative_folder", "") if meta else ""
                else:
                    status = "PENDIENTE CONFIG. DRIVE"
                    archivo = ""
                    carpeta = ""
                resultados.append({
                    "ZONA": zona,
                    "N° POSTE": poste,
                    "TIPO": tipo,
                    "ARCHIVO ESPERADO": f"{poste}_{'A' if tipo == 'ANTES' else 'D'}.*",
                    "ARCHIVO ENCONTRADO": archivo,
                    "CARPETA": carpeta,
                    "RESULTADO": status,
                })
        if resultados:
            qc = pd.DataFrame(resultados)
            st.dataframe(qc, use_container_width=True, hide_index=True)
            errores = qc[qc["RESULTADO"] != "OK"]
            st.caption(f"Evidencias revisadas: {len(qc)} · Requieren atención: {len(errores)}")

# -----------------------------------------------------------------------------
# TABLA COMPLETA
# -----------------------------------------------------------------------------
def preparar_tabla_detalle_con_links(data_df: pd.DataFrame):
    """
    Convierte FOTO ANTES / FOTO DESPUES en enlaces clicables a Google Drive.

    El texto visible sigue siendo el código de la evidencia (por ejemplo 116_A),
    pero al hacer clic se abre la fotografía correspondiente en una pestaña nueva.
    La resolución usa la misma lógica segura ZONA + N° POSTE + A/D del visor.
    """
    tabla = data_df.copy()
    column_config = {}

    if "N° POSTE" not in tabla.columns:
        return tabla, column_config

    links_antes = []
    links_despues = []

    for _, row in tabla.iterrows():
        zona = safe_text(row.get("ZONA", ""))
        poste = normalize_poste_code(row.get("N° POSTE", ""))

        ref_a = safe_text(row.get("FOTO ANTES", "")) if "FOTO ANTES" in tabla.columns else ""
        ref_d = safe_text(row.get("FOTO DESPUES", "")) if "FOTO DESPUES" in tabla.columns else ""

        meta_a = None
        meta_d = None

        if drive_files_antes:
            meta_a, _ = resolver_foto_poste(
                zona, poste, "ANTES", drive_files_antes, explicit_ref=ref_a
            )
        if drive_files_despues:
            meta_d, _ = resolver_foto_poste(
                zona, poste, "DESPUES", drive_files_despues, explicit_ref=ref_d
            )

        if meta_a and meta_a.get("id"):
            nombre_a = Path(str(meta_a.get("name", f"{poste}_A"))).stem
            links_antes.append(f"{drive_view_url(meta_a['id'])}#{nombre_a}")
        else:
            links_antes.append("")

        if meta_d and meta_d.get("id"):
            nombre_d = Path(str(meta_d.get("name", f"{poste}_D"))).stem
            links_despues.append(f"{drive_view_url(meta_d['id'])}#{nombre_d}")
        else:
            links_despues.append("")

    # Se crean las columnas aunque la hoja original no tenga referencias; así, si
    # la evidencia existe en Drive por convención, el enlace aparece igualmente.
    tabla["FOTO ANTES"] = links_antes
    tabla["FOTO DESPUES"] = links_despues

    column_config = {
        "FOTO ANTES": st.column_config.LinkColumn(
            "FOTO ANTES",
            help="Haz clic en el código para abrir la fotografía ANTES en Google Drive.",
            display_text=r"#([^#]+)$",
            width="small",
        ),
        "FOTO DESPUES": st.column_config.LinkColumn(
            "FOTO DESPUÉS",
            help="Haz clic en el código para abrir la fotografía DESPUÉS en Google Drive.",
            display_text=r"#([^#]+)$",
            width="small",
        ),
    }
    return tabla, column_config


st.subheader(f"📋 Detalle de registros filtrados ({len(df_filtrado)} elementos)")
df_detalle_links, detalle_column_config = preparar_tabla_detalle_con_links(df_filtrado)
st.dataframe(
    df_detalle_links.style.apply(resaltar_filas, axis=1),
    use_container_width=True,
    column_config=detalle_column_config,
)

# -----------------------------------------------------------------------------
# EXPORTACIONES
# -----------------------------------------------------------------------------
st.markdown("### 📥 Exportar reportes de obra")
col_exp1, col_exp2 = st.columns(2)

with col_exp1:
    try:
        excel_data = generar_excel_estilizado(df_filtrado)
        st.download_button(
            label="📊 Descargar reporte Excel (.xlsx)",
            data=excel_data,
            file_name="reporte_control_postes.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
    except Exception as e:
        st.error(f"Error al generar Excel: {e}")

with col_exp2:
    incluir_fotos_pdf = st.checkbox(
        "Incluir anexo fotográfico en PDF (máx. 20 postes)",
        value=False,
        help="Puede aumentar el tiempo y tamaño del PDF. Requiere Google Drive configurado.",
    )
    try:
        # Para el anexo se combina el índice de antes y después sin duplicados por ID.
        merged_index = {x.get("id"): x for x in (drive_files_antes + drive_files_despues) if x.get("id")}
        pdf_data = generar_pdf_ejecutivo(
            df_filtrado,
            total_postes,
            pendientes,
            atendidos,
            conformes,
            indexed_files=list(merged_index.values()),
            incluir_anexo_fotos=incluir_fotos_pdf,
            max_fotos=20,
        )
        st.download_button(
            label="📄 Descargar reporte ejecutivo PDF",
            data=pdf_data,
            file_name="reporte_ejecutivo_postes.pdf",
            mime="application/pdf",
            use_container_width=True,
        )
    except Exception as e:
        st.error(f"Error al generar PDF: {e}")

st.markdown("---")
st.caption(
    "Google Sheets: caché de 60 s · Google Drive: caché de 5 min. "
    "Las fotos se resuelven automáticamente por ZONA + N° POSTE + sufijo A/D. "
    "Usa el botón lateral para forzar una actualización inmediata."
)
